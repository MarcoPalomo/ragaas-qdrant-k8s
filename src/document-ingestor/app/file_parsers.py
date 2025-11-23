"""File parsers for different document types."""

import io
from abc import ABC, abstractmethod
from typing import Optional

from pydantic import BaseModel


class ParsedDocument(BaseModel):
    """Parsed document with content and metadata."""

    content: str
    metadata: dict
    filename: str
    file_type: str


class BaseParser(ABC):
    """Base class for file parsers."""

    @abstractmethod
    def can_parse(self, filename: str, content_type: Optional[str] = None) -> bool:
        """Check if this parser can handle the file."""
        pass

    @abstractmethod
    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse the file content."""
        pass


class TextParser(BaseParser):
    """Parser for plain text files."""

    EXTENSIONS = {".txt", ".md", ".rst", ".csv", ".log", ".json", ".xml", ".yaml", ".yml"}

    def can_parse(self, filename: str, content_type: Optional[str] = None) -> bool:
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        return ext in self.EXTENSIONS or (content_type and "text" in content_type)

    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        text = content.decode("utf-8", errors="ignore")
        return ParsedDocument(
            content=text,
            metadata={
                "char_count": len(text),
                "line_count": text.count("\n") + 1,
            },
            filename=filename,
            file_type="text",
        )


class PDFParser(BaseParser):
    """Parser for PDF files."""

    def can_parse(self, filename: str, content_type: Optional[str] = None) -> bool:
        return filename.lower().endswith(".pdf") or content_type == "application/pdf"

    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content))
            text_parts = []
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(f"[Page {page_num + 1}]\n{text}")

            full_text = "\n\n".join(text_parts)
            return ParsedDocument(
                content=full_text,
                metadata={
                    "page_count": len(reader.pages),
                    "char_count": len(full_text),
                },
                filename=filename,
                file_type="pdf",
            )
        except ImportError:
            raise RuntimeError("pypdf is required for PDF parsing. Install with: pip install pypdf")


class DocxParser(BaseParser):
    """Parser for Word documents."""

    def can_parse(self, filename: str, content_type: Optional[str] = None) -> bool:
        return filename.lower().endswith(".docx") or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            full_text = "\n\n".join(paragraphs)

            return ParsedDocument(
                content=full_text,
                metadata={
                    "paragraph_count": len(paragraphs),
                    "char_count": len(full_text),
                },
                filename=filename,
                file_type="docx",
            )
        except ImportError:
            raise RuntimeError("python-docx is required for DOCX parsing. Install with: pip install python-docx")


class HTMLParser(BaseParser):
    """Parser for HTML files."""

    def can_parse(self, filename: str, content_type: Optional[str] = None) -> bool:
        return filename.lower().endswith((".html", ".htm")) or (content_type and "html" in content_type)

    async def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(content, "html.parser")

            # Remove script and style elements
            for element in soup(["script", "style", "nav", "footer", "header"]):
                element.decompose()

            text = soup.get_text(separator="\n")
            # Clean up whitespace
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            full_text = "\n".join(lines)

            title = soup.title.string if soup.title else None

            return ParsedDocument(
                content=full_text,
                metadata={
                    "title": title,
                    "char_count": len(full_text),
                },
                filename=filename,
                file_type="html",
            )
        except ImportError:
            raise RuntimeError("beautifulsoup4 is required for HTML parsing. Install with: pip install beautifulsoup4")


class DocumentParserFactory:
    """Factory for creating document parsers."""

    def __init__(self):
        self.parsers: list[BaseParser] = [
            TextParser(),
            PDFParser(),
            DocxParser(),
            HTMLParser(),
        ]

    def get_parser(self, filename: str, content_type: Optional[str] = None) -> Optional[BaseParser]:
        """Get the appropriate parser for a file."""
        for parser in self.parsers:
            if parser.can_parse(filename, content_type):
                return parser
        return None

    async def parse(self, content: bytes, filename: str, content_type: Optional[str] = None) -> ParsedDocument:
        """Parse a document using the appropriate parser."""
        parser = self.get_parser(filename, content_type)
        if parser is None:
            # Default to text parser
            parser = TextParser()
        return await parser.parse(content, filename)


# Global factory instance
parser_factory = DocumentParserFactory()
